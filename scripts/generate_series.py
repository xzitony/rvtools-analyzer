#!/usr/bin/env python3
"""Derive a monthly series of RVTools exports from the sample export, for trend-mode demos and tests.

The newest sample export is treated as the latest snapshot (September); June, July and August are
reconstructed backwards with known changes, so trend mode has something real to find:
  - VMs added in later months (absent earlier) and legacy VMs retired (present earlier only)
  - VMs resized (half the vCPU or memory before their resize month) and disks grown
  - VMs moved from Dev to Prod, DRS host moves, power-state changes, a rename, VM hardware upgrades
  - data growth per VM (earlier in-use = later / (1 + monthly rate)^months)
  - host esx-prod06 added in August, esx-dev03 upgraded from ESXi 7.0.3 to 8.0.2 in August
  - datastore prod-vmfs-05 expanded from 16 TB to 20 TB in August

usage: python3 scripts/generate_series.py [output-dir]      (requires openpyxl)
"""
import datetime as dt
import glob
import os
import random
import sys

from openpyxl import Workbook, load_workbook

R = random.Random(7)
HERE = os.path.dirname(os.path.abspath(__file__))
BASE = sorted(glob.glob(os.path.join(HERE, "..", "samples", "RVTools_export_all_*.xlsx")))[-1]
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "samples", "series")
DATES = [dt.datetime(2026, 6, 1, 9, 30), dt.datetime(2026, 7, 1, 9, 30), dt.datetime(2026, 8, 3, 9, 30), dt.datetime(2026, 9, 1, 9, 30)]
LAST = len(DATES) - 1
MIB_TB = 1024 * 1024

wb = load_workbook(BASE, read_only=True)
base = {}
for ws in wb.worksheets:
    it = ws.iter_rows(values_only=True)
    base[ws.title] = (list(next(it)), [list(r) for r in it])


def ix(tab, name):
    headers = base[tab][0]
    return headers.index(name) if name in headers else None


info_h, info = base["vInfo"]
C = {k: ix("vInfo", k) for k in ["VM", "VM ID", "VM UUID", "Template", "Powerstate", "CPUs", "Memory", "In Use MiB", "Provisioned MiB",
                                  "Total disk capacity MiB", "Cluster", "Host", "HW version", "Creation date", "Disks"]}
vms = [r for r in info if r[C["Template"]] == "False"]
by_id = {r[C["VM ID"]]: r for r in vms}
hosts_by_cluster = {}
for h in base["vHost"][1]:
    hosts_by_cluster.setdefault(h[ix("vHost", "Cluster")] or "", []).append(h[ix("vHost", "Host")])

ids = sorted(by_id)
R.shuffle(ids)
running = [i for i in ids if by_id[i][C["Powerstate"]] == "poweredOn"]
prod = [i for i in running if by_id[i][C["Cluster"]] == "Prod-Cluster"]

added = {vid: R.choice([1, 2, 3]) for vid in ids[:18]}                 # first snapshot the VM exists in
resized = {vid: R.choice([1, 2, 3]) for vid in ids[18:32]}             # snapshot from which today's size applies
last_disk = {}
for r in base["vDisk"][1]:
    last_disk[r[ix("vDisk", "VM ID")]] = r[ix("vDisk", "Capacity MiB")] or 0
disk_grown = {vid: R.choice([1, 2, 3]) for vid in [i for i in ids[32:] if last_disk.get(i, 0) > 200 * 1024][:8]}  # a disk was 100 GB smaller before this
hw_upgraded = {vid: R.choice([1, 2, 3]) for vid in ids[40:60] if (by_id[vid][C["HW version"]] or 0) >= 17}
powered_on = {vid: R.choice([1, 2, 3]) for vid in running[:6] if vid not in added}
moved_from_dev = {vid: R.choice([1, 2, 3]) for vid in [p for p in prod if p not in added][:5]}
renamed = {ids[61]: 2}
growth = {vid: max(0.0, R.gauss(0.015, 0.012)) for vid in ids}         # monthly data growth per VM
ds_growth = {r[ix("vDatastore", "Name")]: R.choice([0.0, 0.005, 0.01, 0.015, 0.02, 0.03]) for r in base["vDatastore"][1]}
retired = [(f"PRD-LEGACY-{n:02d}", R.choice(prod), R.choice([1, 2, 3])) for n in range(1, 9)]  # (name, template VM, first snapshot it's gone)


def mutate_vm_rows(tabs, s):
    """Apply per-VM overrides for snapshot s to every tab that has a VM ID column."""
    k = LAST - s
    overrides = {}
    for vid, r in by_id.items():
        o = {}
        if vid in resized and s < resized[vid]:
            if hash(vid) % 2:
                o["cpus"] = max(1, r[C["CPUs"]] // 2)
            else:
                o["mem"] = max(1024, r[C["Memory"]] // 2)
        if vid in moved_from_dev and s < moved_from_dev[vid]:
            o["cluster"], o["host"] = "Dev-Cluster", R.choice(hosts_by_cluster["Dev-Cluster"])
        elif R.random() < 0.08 and r[C["Cluster"]]:
            o["host"] = R.choice(hosts_by_cluster.get(r[C["Cluster"]], [r[C["Host"]]]))  # DRS move
        if s < 2 and r[C["Host"]] == "esx-prod06.corp.example":
            o["host"] = R.choice([h for h in hosts_by_cluster["Prod-Cluster"] if "prod06" not in h])
        if vid in powered_on and s < powered_on[vid]:
            o["power"] = "poweredOff"
        if vid in renamed and s < renamed[vid]:
            o["name"] = "TEMP-" + r[C["VM"]]
        if vid in hw_upgraded and s < hw_upgraded[vid]:
            o["hw"] = 13
        o["factor"] = (1 + growth[vid]) ** k
        o["disk_delta"] = -100 * 1024 if vid in disk_grown and s < disk_grown[vid] else 0
        overrides[vid] = o

    for tab, (headers, rows) in tabs.items():
        cid = headers.index("VM ID") if "VM ID" in headers else None
        if cid is None:
            continue
        col = {h: i for i, h in enumerate(headers)}
        for r in rows:
            o = overrides.get(r[cid])
            if not o:
                continue
            if "name" in o and "VM" in col: r[col["VM"]] = o["name"]
            if "power" in o and "Powerstate" in col: r[col["Powerstate"]] = o["power"]
            if "cluster" in o and "Cluster" in col: r[col["Cluster"]] = o["cluster"]
            if "host" in o and "Host" in col: r[col["Host"]] = o["host"]
            if tab in ("vInfo", "vCPU") and "cpus" in o: r[col["CPUs"]] = o["cpus"]
            if tab == "vInfo" and "mem" in o: r[col["Memory"]] = o["mem"]
            if tab == "vMemory" and "mem" in o: r[col["Size MiB"]] = o["mem"]
            if tab == "vInfo":
                if "hw" in o: r[col["HW version"]] = o["hw"]
                r[col["In Use MiB"]] = int((r[col["In Use MiB"]] or 0) / o["factor"])
                if o["disk_delta"]:
                    r[col["Provisioned MiB"]] = (r[col["Provisioned MiB"]] or 0) + o["disk_delta"]
                    r[col["Total disk capacity MiB"]] = (r[col["Total disk capacity MiB"]] or 0) + o["disk_delta"]
            if tab == "vTools" and "hw" in o:
                r[col["VM Version"]] = o["hw"]
            if tab == "vPartition":
                cap = r[col["Capacity MiB"]] or 0
                consumed = int((r[col["Consumed MiB"]] or 0) / o["factor"])
                r[col["Consumed MiB"]], r[col["Free MiB"]] = consumed, cap - consumed
                r[col["Free %"]] = int(round((cap - consumed) / cap * 100)) if cap else 0
        if tab == "vDisk":
            seen = set()
            for r in reversed(rows):   # shrink the last disk of VMs that grew a disk later
                o = overrides.get(r[cid])
                if o and o["disk_delta"] and r[cid] not in seen:
                    r[col["Capacity MiB"]] += o["disk_delta"]
                    seen.add(r[cid])


def build(s):
    tabs = {t: (h, [list(r) for r in rows]) for t, (h, rows) in base.items()}
    k = LAST - s
    # Legacy VMs that were retired later: clone every row of a template VM under a new identity.
    for n, (name, src, gone) in enumerate(retired):
        if s >= gone:
            continue
        new_id, new_uuid = f"vm-9{n:03d}", f"4299{n:028d}"
        for tab, (headers, rows) in tabs.items():
            if "VM ID" not in headers:
                continue
            cid, col = headers.index("VM ID"), {h: i for i, h in enumerate(headers)}
            for r in [r for r in base[tab][1] if r[cid] == src]:
                c = list(r)
                c[cid] = new_id
                if "VM UUID" in col: c[col["VM UUID"]] = new_uuid
                if "VM" in col: c[col["VM"]] = name
                rows.append(c)
    mutate_vm_rows(tabs, s)
    # VMs that didn't exist yet, and snapshots taken after this export
    absent = {vid for vid, first in added.items() if s < first}
    for tab, (headers, rows) in tabs.items():
        if "VM ID" in headers:
            cid = headers.index("VM ID")
            rows[:] = [r for r in rows if r[cid] not in absent]
    sh, srows = tabs["vSnapshot"]
    dcol = sh.index("Date / time")
    srows[:] = [r for r in srows if not r[dcol] or r[dcol] <= DATES[s]]
    # Hosts: prod06 arrives in August; dev03 was on ESXi 7.0.3 before August
    if s < 2:
        for tab, (headers, rows) in tabs.items():
            if "Host" in headers and "VM ID" not in headers:
                hcol = headers.index("Host")
                rows[:] = [r for r in rows if r[hcol] != "esx-prod06.corp.example"]
        hh, hrows = tabs["vHost"]
        for r in hrows:
            if r[hh.index("Host")] == "esx-dev03.corp.example":
                r[hh.index("ESX Version")] = "VMware ESXi 7.0.3 build-21930508"
        for r in tabs["vCluster"][1]:
            if r[tabs["vCluster"][0].index("Name")] == "Prod-Cluster":
                r[tabs["vCluster"][0].index("NumHosts")] = 5
    # Datastores: used space grew 0-3%/month (per datastore); prod-vmfs-05 expanded in August
    dh, drows = tabs["vDatastore"]
    dc = {h: i for i, h in enumerate(dh)}
    for r in drows:
        cap = r[dc["Capacity MiB"]]
        if r[dc["Name"]] == "prod-vmfs-05" and s < 2:
            cap = 16 * MIB_TB
            r[dc["Capacity MiB"]] = cap
        used = (r[dc["In Use MiB"]] or 0) / ((1 + ds_growth[r[dc["Name"]]]) ** k)
        r[dc["In Use MiB"]] = int(used)
        r[dc["Free MiB"]] = int(max(cap - used, cap * 0.02))
        r[dc["Free %"]] = int(round(r[dc["Free MiB"]] / cap * 100))
        r[dc["Provisioned MiB"]] = int((r[dc["Provisioned MiB"]] or 0) / (1.02 ** k))
        if s < 2:
            r[dc["Hosts"]] = ", ".join(h for h in (r[dc["Hosts"]] or "").split(", ") if "prod06" not in h)
    # Recount per-host VM figures so the export stays internally consistent
    ih, irows = tabs["vInfo"]
    hh, hrows = tabs["vHost"]
    for r in hrows:
        name = r[hh.index("Host")]
        on = [v for v in irows if v[ih.index("Host")] == name and v[ih.index("Template")] == "False"]
        r[hh.index("# VMs total")] = len(on)
        running_vms = [v for v in on if v[ih.index("Powerstate")] == "poweredOn"]
        r[hh.index("# VMs")] = len(running_vms)
        r[hh.index("# vCPUs")] = sum(v[ih.index("CPUs")] for v in running_vms)
    mh, mrows = tabs["vMetaData"]
    mrows[0][mh.index("xlsx creation datetime")] = DATES[s]
    return tabs


os.makedirs(OUT, exist_ok=True)
for s, when in enumerate(DATES):
    tabs = build(s)
    out = Workbook()
    out.remove(out.active)
    for tab, (headers, rows) in tabs.items():
        ws = out.create_sheet(tab)
        ws.append(headers)
        for r in rows:
            ws.append(r)
    path = os.path.join(OUT, f"RVTools_export_all_{when:%Y-%m-%d_%H.%M.%S}.xlsx")
    out.save(path)
    print(f"wrote {path}  ({len([r for r in tabs['vInfo'][1] if r[C['Template']] == 'False'])} VMs)")
